
"""
MAI Nexus AI Education
Excel import API for student records.
Author: MAI Solution Hub
"""

from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.education_permissions import require_education_permission
from app.database import get_db
from app.models.education_school import EducationSchool
from app.models.education_student import EducationStudent

router = APIRouter(
    prefix="/education/import",
    tags=["Education Import"],
)

HEADERS = [
    "student_number",
    "full_name",
    "school_id",
    "is_active",
]
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 500


def get_tenant_id(current_user: dict) -> int:
    tenant_id = current_user.get("tenant_id")
    if tenant_id is None:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Pengguna belum ditetapkan ke tenant Education",
        )
    return int(tenant_id)


def parse_active(value):
    if value is None or str(value).strip() == "":
        return True

    normalized = str(value).strip().lower()

    if normalized in {"true", "1", "ya", "aktif"}:
        return True

    if normalized in {"false", "0", "tidak", "nonaktif", "non-aktif"}:
        return False

    raise ValueError(
        "is_active harus TRUE/FALSE, 1/0, ya/tidak, atau aktif/nonaktif"
    )


@router.get("/students/template")
def download_student_template(
    current_user: dict = Depends(
        require_education_permission("education:students:create")
    ),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(HEADERS)
    sheet.append(["S-0001", "Nama Contoh", 1, True])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 15
    sheet.column_dimensions["D"].width = 15

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="template_data_siswa.xlsx"'
            )
        },
    )


@router.post("/students")
async def import_students_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: dict = Depends(
        require_education_permission("education:students:create")
    ),
):
    tenant_id = get_tenant_id(current_user)

    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Unggah file Excel berformat .xlsx",
        )

    content = await file.read(MAX_FILE_SIZE + 1)
    await file.close()

    if not content:
        raise HTTPException(status_code=400, detail="File Excel kosong")

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="Ukuran file maksimal 5 MB",
        )

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="File Excel tidak valid atau rusak",
        )

    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)

    if not header_row:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail="File tidak memiliki baris header",
        )

    normalized_headers = [
        str(value).strip().lower() if value is not None else ""
        for value in header_row
    ]

    if normalized_headers != HEADERS:
        workbook.close()
        raise HTTPException(
            status_code=400,
            detail={
                "message": "Header Excel tidak sesuai template",
                "required_headers": HEADERS,
                "received_headers": normalized_headers,
            },
        )

    results = []
    seen_numbers = set()
    processed = 0
    imported = 0
    failed = 0

    try:
        for excel_row, values in enumerate(rows, start=2):
            if not values or all(
                value is None or str(value).strip() == ""
                for value in values
            ):
                continue

            processed += 1

            if processed > MAX_ROWS:
                results.append({
                    "row": excel_row,
                    "student_number": None,
                    "status": "failed",
                    "message": f"Maksimal {MAX_ROWS} baris data per file",
                })
                failed += 1
                break

            student_number = None

            try:
                student_number = (
                    str(values[0]).strip()
                    if len(values) > 0 and values[0] is not None
                    else ""
                )
                full_name = (
                    str(values[1]).strip()
                    if len(values) > 1 and values[1] is not None
                    else ""
                )
                school_value = values[2] if len(values) > 2 else None
                active_value = values[3] if len(values) > 3 else None

                if not student_number:
                    raise ValueError("student_number wajib diisi")

                if len(student_number) > 50:
                    raise ValueError(
                        "student_number maksimal 50 karakter"
                    )

                if not full_name:
                    raise ValueError("full_name wajib diisi")

                if len(full_name) > 150:
                    raise ValueError("full_name maksimal 150 karakter")

                if isinstance(school_value, bool) or school_value is None:
                    raise ValueError("school_id wajib berupa ID sekolah")

                school_number = float(school_value)
                if not school_number.is_integer() or school_number <= 0:
                    raise ValueError("school_id harus berupa bilangan bulat positif")
                school_id = int(school_number)

                is_active = parse_active(active_value)

                if student_number in seen_numbers:
                    raise ValueError(
                        "Nomor siswa duplikat dalam file Excel"
                    )
                seen_numbers.add(student_number)

                school = (
                    db.query(EducationSchool)
                    .filter(
                        EducationSchool.id == school_id,
                        EducationSchool.tenant_id == tenant_id,
                    )
                    .first()
                )

                if school is None:
                    raise ValueError(
                        "Sekolah tidak ditemukan pada tenant ini"
                    )

                duplicate = (
                    db.query(EducationStudent)
                    .filter(
                        EducationStudent.tenant_id == tenant_id,
                        EducationStudent.student_number == student_number,
                    )
                    .first()
                )

                if duplicate:
                    raise ValueError(
                        "Nomor siswa sudah terdaftar pada tenant ini"
                    )

                student = EducationStudent(
                    tenant_id=tenant_id,
                    school_id=school_id,
                    student_number=student_number,
                    full_name=full_name,
                    is_active=is_active,
                )
                db.add(student)
                db.commit()
                imported += 1

                results.append({
                    "row": excel_row,
                    "student_number": student_number,
                    "status": "imported",
                    "message": "Berhasil diimpor",
                })

            except (ValueError, TypeError, IndexError) as exc:
                db.rollback()
                failed += 1
                results.append({
                    "row": excel_row,
                    "student_number": student_number,
                    "status": "failed",
                    "message": str(exc),
                })

            except IntegrityError:
                db.rollback()
                failed += 1
                results.append({
                    "row": excel_row,
                    "student_number": student_number,
                    "status": "failed",
                    "message": "Data bertentangan dengan data yang sudah ada",
                })

            except Exception:
                db.rollback()
                failed += 1
                results.append({
                    "row": excel_row,
                    "student_number": student_number,
                    "status": "failed",
                    "message": "Terjadi kesalahan saat menyimpan data",
                })
    finally:
        workbook.close()

    return {
        "message": "Proses impor selesai",
        "total_processed": processed,
        "imported": imported,
        "failed": failed,
        "results": results,
    }
